import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


def s(v):
    if v is None:
        return ""

    if isinstance(v, list):
        return "; ".join(
            json.dumps(x, ensure_ascii=False)
            if isinstance(x, dict)
            else str(x)
            for x in v
        )

    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)

    return str(v)


def get_next_row(ws):
    row = ws.max_row

    while row > 1:
        if any(
            ws.cell(row, c).value
            for c in range(1, ws.max_column + 1)
        ):
            break
        row -= 1

    return row + 1


def find_application_kit(job_url, applications_dir):
    """
    Find an existing application kit by matching the job URL
    stored inside application_kit.json.
    """

    applications_dir = Path(applications_dir)

    if not applications_dir.exists():
        return None

    for kit_dir in applications_dir.iterdir():

        if not kit_dir.is_dir():
            continue

        kit_file = kit_dir / "application_kit.json"

        if not kit_file.exists():
            continue

        try:
            kit = json.loads(
                kit_file.read_text(
                    encoding="utf-8"
                )
            )
        except (OSError, json.JSONDecodeError):
            continue

        kit_job = kit.get("job") or {}

        kit_url = (
            kit_job.get("url")
            or kit.get("job_url")
            or ""
        )

        if str(kit_url).strip() == str(job_url).strip():
            return kit_dir

    return None


def ensure_apply_assist_columns(ws):
    """
    Automatically add Apply-Assist columns if they do not already exist.

    This makes the Excel CRM schema self-maintaining.
    """

    required_headers = [
        "Application Kit Path",
        "Application Kit Status",
        "Cover Letter Status",
        "Resume Bullets Status",
        "Screening Answers Status",
        "Apply Status",
    ]

    existing_headers = {
        ws.cell(1, c).value
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }

    next_column = ws.max_column + 1

    added = []

    for header in required_headers:

        if header not in existing_headers:

            ws.cell(
                1,
                next_column
            ).value = header

            existing_headers.add(header)

            added.append(header)

            next_column += 1

    return added


def main():

    ap = argparse.ArgumentParser()

    ap.add_argument(
        "--analysis",
        required=True
    )

    ap.add_argument(
        "--crm",
        required=True
    )

    ap.add_argument(
        "--applications-dir",
        default="applications"
    )

    args = ap.parse_args()

    # --------------------------------------------------
    # LOAD ANALYSIS
    # --------------------------------------------------

    data = json.loads(
        Path(args.analysis).read_text(
            encoding="utf-8"
        )
    )

    # Support both:
    # { ... }
    # and
    # [ {...}, {...} ]

    if isinstance(data, dict):
        data = [data]

    if not isinstance(data, list):
        raise ValueError(
            "Analysis file must contain an object or list of objects."
        )

    # --------------------------------------------------
    # LOAD WORKBOOK
    # --------------------------------------------------

    wb = load_workbook(args.crm)

    if "Jobs" not in wb.sheetnames:
        raise ValueError(
            "Could not find 'Jobs' sheet in CRM workbook."
        )

    ws = wb["Jobs"]

    # --------------------------------------------------
    # ENSURE APPLY-ASSIST COLUMNS
    # --------------------------------------------------

    added_headers = ensure_apply_assist_columns(ws)

    if added_headers:

        print()
        print("Apply-Assist CRM columns added:")

        for header in added_headers:
            print(f"  + {header}")

    else:

        print()
        print("✓ Apply-Assist CRM columns already exist")

    # --------------------------------------------------
    # READ HEADERS
    # --------------------------------------------------

    headers = {
        ws.cell(1, c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }

    # --------------------------------------------------
    # FIND EXISTING JOBS BY URL
    # --------------------------------------------------

    url_col = headers.get("Job URL")

    if not url_col:
        raise ValueError(
            "Could not find 'Job URL' column in Jobs sheet."
        )

    existing_by_url = {}

    for row in range(2, ws.max_row + 1):

        value = ws.cell(row, url_col).value

        if value:
            existing_by_url[
                str(value).strip()
            ] = row

    # --------------------------------------------------
    # PROCESS
    # --------------------------------------------------

    imported = 0
    updated = 0
    skipped = 0
    kits_found = 0

    target = get_next_row(ws)

    for item in data:

        if not isinstance(item, dict):
            continue

        company = s(item.get("company"))
        role = s(item.get("role"))
        job_url = s(item.get("job_url"))

        if not job_url:
            skipped += 1
            continue

        # --------------------------------------------------
        # NESTED DATA
        # --------------------------------------------------

        stipend = item.get("stipend") or {}
        ppo = item.get("ppo") or {}
        eligibility = item.get("eligibility") or {}
        availability = item.get("availability") or {}
        scores = item.get("scores") or {}
        reasoning = item.get("reasoning") or {}

        overall_score = scores.get("overall_score")

        # --------------------------------------------------
        # FIND APPLY-ASSIST KIT
        # --------------------------------------------------

        kit_dir = find_application_kit(
            job_url,
            args.applications_dir
        )

        kit_exists = bool(kit_dir)

        if kit_exists:
            kits_found += 1

        kit_path = (
            str(kit_dir.resolve())
            if kit_dir
            else ""
        )

        # --------------------------------------------------
        # DETERMINE ROW
        # --------------------------------------------------

        if job_url in existing_by_url:

            row_number = existing_by_url[job_url]

            updated += 1

            action = "Updated"

        else:

            row_number = target

            target += 1

            imported += 1

            job_id = (
                f"{company[:12].upper().replace(' ', '-')}"
                f"-{date.today():%y%m%d}"
                f"-{imported}"
            )

            if "Job ID" in headers:

                ws.cell(
                    row_number,
                    headers["Job ID"]
                ).value = job_id

            action = "Imported"

        # --------------------------------------------------
        # VALUES TO WRITE
        # --------------------------------------------------

        values = {

            "Company":
                company,

            "Role":
                role,

            "Job URL":
                job_url,

            "Source":
                "AI Analyzer",

            "Location":
                s(item.get("location")),

            "Work Mode":
                s(item.get("work_mode")),

            "Job Type":
                s(item.get("job_type")),

            "Stipend":
                s(stipend.get("raw")),

            "Duration":
                s(item.get("duration")),

            "PPO Signal":
                s(ppo.get("status")),

            "Required Skills":
                s(item.get("required_skills")),

            "Preferred Skills":
                s(item.get("preferred_skills")),

            "JD Text":
                s(
                    (item.get("_job") or {}).get(
                        "snippet"
                    )
                ),

            "Match Score":
                overall_score,

            "Technical Fit":
                scores.get("technical_fit"),

            "Experience Fit":
                scores.get("experience_fit"),

            "Role Fit":
                scores.get("role_fit"),

            "PPO Score":
                scores.get("ppo_score"),

            "Compensation Fit":
                scores.get("compensation_fit"),

            "Opportunity Quality":
                scores.get("opportunity_quality"),

            "Priority":
                s(item.get("recommendation")),

            "Recommendation":
                s(item.get("recommendation")),

            "Skill Gaps":
                s(item.get("skill_gaps")),

            "Resume Version":
                s(item.get("recommended_resume")),

            "Date Found":
                date.today().isoformat(),

            "Notes":
                s(
                    reasoning.get("why_apply") or []
                ),

            "Eligibility Status":
                s(eligibility.get("status")),

            "Eligibility Reasons":
                s(eligibility.get("reasons")),

            "Availability Status":
                s(availability.get("status")),

            "Availability Reason":
                s(availability.get("reason")),

            "Compensation Status":
                s(stipend.get("status")),

            "Guaranteed Stipend":
                stipend.get("guaranteed_min_inr"),

            "Incentive Included":
                (
                    "Yes"
                    if stipend.get(
                        "incentive_included"
                    )
                    else "No"
                ),

            "PPO Status":
                s(ppo.get("status")),

            "PPO Evidence":
                s(ppo.get("evidence")),

            "Hard Blockers":
                s(item.get("hard_blockers")),

            "Red Flags":
                s(item.get("red_flags")),

            "Interview Topics":
                s(item.get("interview_topics")),
        }

        # --------------------------------------------------
        # APPLY-ASSIST VALUES
        # --------------------------------------------------

        values.update({

            "Application Kit Path":
                kit_path,

            "Application Kit Status":
                (
                    "Ready for Review"
                    if kit_exists
                    else ""
                ),

            "Cover Letter Status":
                (
                    "Generated"
                    if kit_exists
                    else ""
                ),

            "Resume Bullets Status":
                (
                    "Generated"
                    if kit_exists
                    else ""
                ),

            "Screening Answers Status":
                (
                    "Generated"
                    if kit_exists
                    else ""
                ),

            "Apply Status":
                (
                    "Ready to Apply"
                    if (
                        kit_exists
                        and s(
                            item.get(
                                "recommendation"
                            )
                        ).upper()
                        in {
                            "APPLY",
                            "APPLY ASAP"
                        }
                    )
                    else ""
                ),
        })

        # --------------------------------------------------
        # WRITE VALUES
        # --------------------------------------------------

        for header, value in values.items():

            if header in headers:

                ws.cell(
                    row_number,
                    headers[header]
                ).value = value

        existing_by_url[job_url] = row_number

        kit_note = (
            " | Kit: Ready"
            if kit_exists
            else ""
        )

        print(
            f"{action}: {company} — {role}"
            f" | Score: {overall_score}"
            f"{kit_note}"
        )

    # --------------------------------------------------
    # SAVE
    # --------------------------------------------------

    try:

        wb.save(args.crm)

    except PermissionError:

        print(
            "\n[!] Permission Error:"
            f" '{args.crm}' is currently open."
        )

        print(
            "Close Excel and run the command again."
        )

        return

    # --------------------------------------------------
    # SUMMARY
    # --------------------------------------------------

    print()
    print("===================================")
    print(f"New jobs imported      : {imported}")
    print(f"Existing jobs updated  : {updated}")
    print(f"Skipped                : {skipped}")
    print(f"Analyses processed     : {len(data)}")
    print(f"Application kits found : {kits_found}")
    print(f"Saved → {args.crm}")
    print("===================================")


if __name__ == "__main__":
    main()
