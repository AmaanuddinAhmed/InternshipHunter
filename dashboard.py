import argparse
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


KEEP = {"APPLY ASAP", "APPLY", "CONSIDER"}

APPLY_STATUSES = [
    "NOT APPLIED",
    "APPLIED",
    "REJECTED",
    "WITHDRAWN",
]


def real_rows(ws):
    """
    Yield each data row's values (min_row=2 onward) for rows that hold at
    least one genuine value.

    Some sheets (Applications, Interviews) are pre-formatted with a
    formula in one column on every row — e.g. Applications' "Days Since
    Applied" is =IF(G2="","",TODAY()-G2) on every row, even ones with no
    application data yet. That formula cell always has a non-empty
    .value (the formula string itself, when the workbook is loaded
    without data_only), so a plain "any value present" check treats
    every pre-formatted blank row as real. We only look at cells that
    aren't themselves formulas when deciding whether a row counts.
    """
    for row in ws.iter_rows(min_row=2):
        has_real_value = any(
            cell.data_type != "f" and cell.value not in (None, "")
            for cell in row
        )
        if has_real_value:
            yield tuple(cell.value for cell in row)

def build_dashboard(path):

    wb = load_workbook(path)

    jobs = wb["Jobs"]
    apps = wb["Applications"]

    headers = {
        str(c.value).strip(): i
        for i, c in enumerate(jobs[1], start=0)
        if c.value
    }

    app_headers = {
        str(c.value).strip(): i
        for i, c in enumerate(apps[1], start=0)
        if c.value
    }

    def g(row, name):
        i = headers.get(name)

        if i is not None and i < len(row):
            return row[i]

        return None

    # ========================================================
    # Ensure Apply-Assist columns exist
    # ========================================================

    required_columns = [
        "Application Kit Path",
        "Application Kit Status",
        "Cover Letter Status",
        "Resume Bullets Status",
        "Screening Answers Status",
        "Apply Status",
    ]

    for column in required_columns:

        if column not in headers:

            new_col = jobs.max_column + 1

            jobs.cell(
                1,
                new_col,
                column
            )

            headers[column] = new_col - 1

    # ========================================================
    # Apply Status dropdown
    # ========================================================

    apply_status_col = headers["Apply Status"] + 1

    validation = DataValidation(
        type="list",
        formula1='"NOT APPLIED,APPLIED,REJECTED,WITHDRAWN"',
        allow_blank=True
    )

    validation.error = (
        "Choose a status from the dropdown."
    )

    validation.errorTitle = (
        "Invalid Apply Status"
    )

    validation.prompt = (
        "Select the current application status."
    )

    validation.promptTitle = (
        "Application Status"
    )

    jobs.add_data_validation(
        validation
    )

    validation.add(
        f"{get_column_letter(apply_status_col)}2:"
        f"{get_column_letter(apply_status_col)}10000"
    )

    # ========================================================
    # Determine already-applied jobs
    # ========================================================

    applied_ids = set()

    if "Job ID" in app_headers:

        for row in real_rows(apps):

            jid = row[
                app_headers["Job ID"]
            ]

            if jid:
                applied_ids.add(
                    str(jid).strip()
                )

    # ========================================================
    # Build job records
    # ========================================================

    records = []

    for row in real_rows(jobs):

        jid = g(row, "Job ID")

        recommendation = str(
            g(row, "Recommendation") or ""
        ).strip().upper()

        if jid is None:
            continue

        if recommendation not in KEEP:
            continue

        score = g(row, "Match Score")

        try:
            score_num = float(score)
        except (TypeError, ValueError):
            score_num = -1

        job_id = str(jid).strip()

        apply_status = str(
            g(row, "Apply Status") or ""
        ).strip().upper()

        # ----------------------------------------------------
        # Application state
        # ----------------------------------------------------

        applied = (
            job_id in applied_ids
            or apply_status == "APPLIED"
        )

        records.append(
            {
                "job_id": job_id,

                "company":
                    g(row, "Company") or "",

                "role":
                    g(row, "Role") or "",

                "url":
                    g(row, "Job URL") or "",

                "location":
                    g(row, "Location") or "",

                "work_mode":
                    g(row, "Work Mode") or "",

                "stipend":
                    g(row, "Stipend") or "",

                "ppo":
                    (
                        g(row, "PPO Status")
                        or g(row, "PPO Signal")
                        or ""
                    ),

                "score":
                    score_num,

                "priority":
                    g(row, "Priority") or "",

                "recommendation":
                    recommendation,

                "resume":
                    g(row, "Resume Version") or "",

                "date_found":
                    g(row, "Date Found") or "",

                "kit_path":
                    g(row, "Application Kit Path")
                    or "",

                "kit_status":
                    g(row, "Application Kit Status")
                    or "",

                "apply_status":
                    apply_status
                    or "NOT APPLIED",

                "applied":
                    applied,
            }
        )

    # ========================================================
    # Sort
    # ========================================================

    records.sort(
        key=lambda x: (
            x["applied"],
            0
            if x["recommendation"] == "APPLY ASAP"
            else 1,
            -x["score"],
        )
    )

    # ========================================================
    # Recreate Morning Dashboard
    # ========================================================

    if "Morning Dashboard" in wb.sheetnames:
        del wb["Morning Dashboard"]

    ws = wb.create_sheet(
        "Morning Dashboard",
        0
    )

    # ========================================================
    # Title
    # ========================================================

    ws["A1"] = (
        "AMAAN INTERNSHIP ECOSYSTEM — MORNING DASHBOARD"
    )

    ws["A1"].font = Font(
        size=18,
        bold=True
    )

    ws.merge_cells("A1:L1")

    ws["A2"] = (
        f"Updated: {date.today().isoformat()}"
    )

    ws["A3"] = (
        "Workflow: Discover → Analyze → Shortlist → "
        "Apply-Assist → Apply → Track"
    )

    ws.merge_cells("A3:L3")

    # ========================================================
    # Statistics
    # ========================================================

    total = len(
        list(real_rows(jobs))
    )

    active_records = [
        r for r in records
        if not r["applied"]
    ]

    applied_records = [
        r for r in records
        if r["applied"]
    ]

    shortlist = len(
        active_records
    )

    asap = sum(
        1
        for r in active_records
        if r["recommendation"] == "APPLY ASAP"
    )

    apply_count = sum(
        1
        for r in active_records
        if r["recommendation"] == "APPLY"
    )

    kits_ready = sum(
        1
        for r in active_records
        if str(r["kit_status"]).strip().upper()
        == "READY"
        or r["kit_path"]
    )

    stats = [
        ("Relevant/Tracked Jobs", total),
        ("Apply Queue", shortlist),
        ("APPLY ASAP", asap),
        ("APPLY", apply_count),
        ("Kits Ready", kits_ready),
        ("Applied", len(applied_records)),
    ]

    for col, (label, value) in enumerate(
        stats,
        start=1
    ):

        ws.cell(
            5,
            col,
            label
        ).font = Font(
            bold=True
        )

        ws.cell(
            6,
            col,
            value
        ).font = Font(
            size=14,
            bold=True
        )

    # ========================================================
    # Apply Queue heading
    # ========================================================

    start = 9

    ws.cell(
        start,
        1,
        "APPLY QUEUE"
    ).font = Font(
        size=14,
        bold=True
    )

    ws.merge_cells(
        start_row=start,
        start_column=1,
        end_row=start,
        end_column=12
    )

    ws.cell(
        start + 1,
        1,
        "Review APPLY ASAP first. Open the job, review the generated kit, then apply manually."
    )

    ws.merge_cells(
        start_row=start + 1,
        start_column=1,
        end_row=start + 1,
        end_column=12
    )

    # ========================================================
    # Main table
    # ========================================================

    header_row = start + 3

    cols = [
        "Rank",
        "Company",
        "Role",
        "Score",
        "Recommendation",
        "Location",
        "Work Mode",
        "Stipend",
        "PPO",
        "Resume",
        "Application Kit",
        "Status",
    ]

    for c, name in enumerate(
        cols,
        start=1
    ):

        cell = ws.cell(
            header_row,
            c,
            name
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    thin = Side(
        style="thin"
    )

    # ========================================================
    # Helper: resolve kit
    # ========================================================

    def resolve_kit_path(kit_path):

        if not kit_path:
            return None

        kit_file = Path(
            str(kit_path).strip()
        )

        if not kit_file.is_absolute():

            kit_file = (
                path.parent / kit_file
            )

        if kit_file.exists():
            return kit_file.resolve()

        return None

    # ========================================================
    # Active Apply Queue
    # ========================================================

    row_number = header_row + 1
    active_rank = 1

    for item in active_records:

        values = [
            active_rank,
            item["company"],
            item["role"],
            (
                item["score"]
                if item["score"] >= 0
                else ""
            ),
            item["recommendation"],
            item["location"],
            item["work_mode"],
            item["stipend"],
            item["ppo"],
            item["resume"],
            "",
            "READY TO APPLY",
        ]

        for c, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row_number,
                c,
                value
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

        # ----------------------------------------------------
        # Job link
        # ----------------------------------------------------

        if item["url"]:

            role_cell = ws.cell(
                row_number,
                3
            )

            role_cell.hyperlink = item["url"]
            role_cell.style = "Hyperlink"

        # ----------------------------------------------------
        # Kit link
        # ----------------------------------------------------

        kit_file = resolve_kit_path(
            item["kit_path"]
        )

        kit_cell = ws.cell(
            row_number,
            11
        )

        if kit_file:

            kit_cell.value = "OPEN KIT"

            kit_cell.hyperlink = (
                kit_file.as_uri()
            )

            kit_cell.style = "Hyperlink"

        elif item["kit_status"]:

            kit_cell.value = (
                item["kit_status"]
            )

        else:

            kit_cell.value = (
                "NOT GENERATED"
            )

        # ----------------------------------------------------
        # Recommendation styling
        # ----------------------------------------------------

        if item["recommendation"] == "APPLY ASAP":

            fill = PatternFill(
                fill_type="solid",
                fgColor="FFF2CC"
            )

        elif item["recommendation"] == "APPLY":

            fill = PatternFill(
                fill_type="solid",
                fgColor="E2F0D9"
            )

        else:

            fill = PatternFill(
                fill_type="solid",
                fgColor="FCE4D6"
            )

        for c in range(
            1,
            len(cols) + 1
        ):

            ws.cell(
                row_number,
                c
            ).fill = fill

            ws.cell(
                row_number,
                c
            ).border = Border(
                bottom=thin
            )

        row_number += 1
        active_rank += 1

    # ========================================================
    # Already Applied
    # ========================================================

    applied_start = row_number + 2

    ws.cell(
        applied_start,
        1,
        "ALREADY APPLIED"
    ).font = Font(
        size=14,
        bold=True
    )

    ws.merge_cells(
        start_row=applied_start,
        start_column=1,
        end_row=applied_start,
        end_column=12
    )

    applied_header = applied_start + 2

    for c, name in enumerate(
        cols,
        start=1
    ):

        cell = ws.cell(
            applied_header,
            c,
            name
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    row_number = applied_header + 1
    applied_rank = 1

    for item in applied_records:

        values = [
            applied_rank,
            item["company"],
            item["role"],
            (
                item["score"]
                if item["score"] >= 0
                else ""
            ),
            item["recommendation"],
            item["location"],
            item["work_mode"],
            item["stipend"],
            item["ppo"],
            item["resume"],
            "",
            "APPLIED",
        ]

        for c, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row_number,
                c,
                value
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

        if item["url"]:

            role_cell = ws.cell(
                row_number,
                3
            )

            role_cell.hyperlink = item["url"]
            role_cell.style = "Hyperlink"

        kit_file = resolve_kit_path(
            item["kit_path"]
        )

        if kit_file:

            kit_cell = ws.cell(
                row_number,
                11
            )

            kit_cell.value = "OPEN KIT"

            kit_cell.hyperlink = (
                kit_file.as_uri()
            )

            kit_cell.style = "Hyperlink"

        for c in range(
            1,
            len(cols) + 1
        ):

            ws.cell(
                row_number,
                c
            ).border = Border(
                bottom=thin
            )

        row_number += 1
        applied_rank += 1

    # ========================================================
    # Instructions
    # ========================================================

    help_row = row_number + 2

    ws.cell(
        help_row,
        1,
        "MORNING ROUTINE"
    ).font = Font(
        bold=True,
        size=12
    )

    steps = [
        "1. Start with APPLY ASAP.",
        "2. Click the Role to open the original job.",
        "3. Click OPEN KIT to review the generated application material.",
        "4. Submit the application manually.",
        "5. In the Jobs sheet, change Apply Status to APPLIED.",
        "6. Refresh the dashboard. The job will move to ALREADY APPLIED.",
    ]

    for i, text in enumerate(
        steps,
        start=help_row + 1
    ):

        ws.cell(
            i,
            1,
            text
        )

        ws.merge_cells(
            start_row=i,
            start_column=1,
            end_row=i,
            end_column=12
        )

    # ========================================================
    # Column widths
    # ========================================================

    widths = [
        7,
        28,
        42,
        9,
        17,
        18,
        14,
        20,
        18,
        36,
        20,
        18,
    ]

    for i, width in enumerate(
        widths,
        start=1
    ):

        ws.column_dimensions[
            get_column_letter(i)
        ].width = width

    # ========================================================
    # Freeze / filter
    # ========================================================

    ws.freeze_panes = (
        f"A{header_row + 1}"
    )

    last_table_row = max(
        header_row,
        row_number - 1
    )

    ws.auto_filter.ref = (
        f"A{header_row}:L{last_table_row}"
    )

    # ========================================================
    # Analytics
    # ========================================================

    if "Analytics" in wb.sheetnames:

        analytics = wb["Analytics"]

        analytics["B2"] = total
        analytics["B3"] = len(records)
        analytics["B4"] = len(
            list(real_rows(apps))
        )

        analytics["B6"] = 0
        analytics["B7"] = 0

    # ========================================================
    # Save
    # ========================================================

    wb.save(path)

    print(
        f"Dashboard refreshed → {path}"
    )

    print(
        f"Apply Queue: {shortlist} "
        f"| APPLY ASAP: {asap} "
        f"| APPLY: {apply_count} "
        f"| Kits Ready: {kits_ready} "
        f"| Applied: {len(applied_records)}"
    )


if __name__ == "__main__":

    ap = argparse.ArgumentParser()

    ap.add_argument(
        "--crm",
        default="Amaan_Internship_Ecosystem_v2.xlsx"
    )

    args = ap.parse_args()

    build_dashboard(
        Path(args.crm)
    )